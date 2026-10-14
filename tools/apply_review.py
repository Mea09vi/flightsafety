"""นำผลจากไฟล์ตรวจ (make_review.py) กลับเข้า data/intl_seed.csv

ใช้:  python tools/apply_review.py review/ชื่อไฟล์.xlsx [--write] [--date YYYY-MM-DD]

ค่าเริ่มต้นแค่แสดงว่าจะเปลี่ยนอะไร ใส่ --write เพื่อเขียนจริง
  ผ่าน  → status=published, reviewed=วันที่ (ถ้าเคสเปลี่ยนสถานะหรือมีการแก้ข้อความ)
  รอแก้ → status=draft หมายเหตุไปอยู่ใน review/todo_<วันที่>.md
  ตัด   → status=cut (แถวยังอยู่ใน seed แต่ build ไม่นำขึ้นแอป)
  ว่าง  → ข้าม

เทียบค่าในไฟล์ตรวจกับ fingerprint ในคอลัมน์ซ่อน _base:
  เซลล์ไม่เปลี่ยน → ไม่แตะ seed
  เซลล์เปลี่ยน และ seed ยังเหมือนตอนสร้างไฟล์ → นำค่าใหม่เข้า
  เซลล์เปลี่ยน และ seed ก็เปลี่ยนไปแล้ว → ชนกัน ข้ามทั้งแถวแล้วรายงาน
ก่อนเขียนจะสำรอง seed ไว้ที่ review/backup/
"""
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from build_data import PII_PATTERNS, TAX, read_csv, split_lines, split_tags
from review_common import (COLUMNS, DECISIONS, EDITABLE, REVIEW_DIR, ROOT, SEED,
                           SHEET, STATUS_OF, fingerprint, norm, read_seed, write_seed)

REQUIRED_TEXT = ["aircraft", "title", "summary", "factors", "actions_today", "brief_question"]


def short(text, start=0, n=60):
    text = text.replace("\n", " / ")
    out = ("…" if start else "") + text[start:start + n]
    return out + ("…" if len(text) > start + n else "")


def diff_pair(old, new):
    """แสดงช่วงที่เริ่มต่างกัน (ข้อความยาวมักแก้แค่ท้าย)"""
    i = next((i for i, (a, b) in enumerate(zip(old, new)) if a != b), min(len(old), len(new)))
    start = max(0, i - 20)
    return short(old, start), short(new, start)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    path = Path(args[0]).resolve()
    write = "--write" in sys.argv
    today = date.today().isoformat()
    if "--date" in sys.argv:
        today = sys.argv[sys.argv.index("--date") + 1]
        date.fromisoformat(today)

    valid = {
        "aircraft": {r["code"] for r in read_csv(TAX / "aircraft.csv")} | {"ALL"},
        "hfacs": {r["code"] for r in read_csv(TAX / "hfacs.csv")},
    }
    fieldnames, seed = read_seed()
    by_id = {r["id"]: r for r in seed}

    wb = load_workbook(path)
    if SHEET not in wb.sheetnames:
        sys.exit(f"ไม่พบแท็บ \"{SHEET}\" ใน {path}")
    ws = wb[SHEET]
    head_to_key = {head: key for key, head, _, _ in COLUMNS}
    col = {}
    for cell in ws[1]:
        key = head_to_key.get(norm(cell.value))
        if key:
            col[key] = cell.column - 1
    missing = [k for k in ["id", "decision", "note", "_base"] if k not in col]
    if missing:
        sys.exit(f"ไม่พบคอลัมน์ {missing} — หัวคอลัมน์ในแท็บ \"{SHEET}\" ถูกแก้หรือเปล่า")
    for k in EDITABLE:
        if k not in col:
            print(f"เตือน: ไม่พบคอลัมน์ {k} จะถือว่าไม่ได้แก้")

    errors, warnings, conflicts, notes = [], [], [], []
    updates = {}   # id → {field: new value}
    count = {d: 0 for d in DECISIONS}
    count["ว่าง"] = 0
    seen = set()

    for cells in ws.iter_rows(min_row=2, values_only=True):
        get = lambda k: norm(cells[col[k]]) if k in col and col[k] < len(cells) else ""
        cid = get("id")
        if not cid:
            continue
        if cid in seen:
            errors.append(f"{cid}: อยู่ในไฟล์ตรวจมากกว่า 1 แถว")
            continue
        seen.add(cid)
        row = by_id.get(cid)
        if row is None:
            errors.append(f"{cid}: ไม่มีใน seed")
            continue
        decision, note = get("decision"), get("note")
        base = get("_base").split(",")
        if len(base) != len(EDITABLE):
            errors.append(f"{cid}: คอลัมน์ _base เสีย (ถูกแก้หรือลบ) — ข้ามแถว")
            continue

        changed, clash = {}, []
        for k, b in zip(EDITABLE, base):
            if k not in col:
                continue
            new = get(k)
            if fingerprint(new) == b:
                continue                       # FSO ไม่ได้แก้
            if norm(row[k]) == new:
                continue                       # seed ถูกแก้เป็นค่าเดียวกันแล้ว
            if fingerprint(row[k]) == b:
                changed[k] = new
            else:
                clash.append(k)

        if decision not in DECISIONS:
            if decision:
                errors.append(f"{cid}: การตัดสิน \"{decision}\" ไม่ใช่ {'/'.join(DECISIONS)}")
                continue
            count["ว่าง"] += 1
            if changed:
                warnings.append(f"{cid}: แก้ {', '.join(changed)} แต่ช่องการตัดสินว่าง — ไม่นำเข้า")
            continue
        count[decision] += 1
        if clash:
            conflicts.append(f"{cid}: {', '.join(clash)} ถูกแก้ทั้งในไฟล์ตรวจและใน seed หลังสร้างไฟล์ — ข้ามทั้งแถว")
            continue

        merged = {**row, **changed}
        bad = [k for k in REQUIRED_TEXT if not merged[k].strip()]
        if bad:
            errors.append(f"{cid}: ช่อง {', '.join(bad)} ว่าง — ข้ามแถว")
            continue
        bad_tags = [f"{k} '{t}'" for k in ("aircraft", "hfacs") for t in split_tags(merged[k])
                    if t not in valid[k]]
        if bad_tags:
            errors.append(f"{cid}: {', '.join(bad_tags)} ไม่อยู่ใน taxonomy — ข้ามแถว")
            continue

        for k, v in changed.items():
            for pattern, what in PII_PATTERNS:
                if pattern.search(v):
                    warnings.append(f"{cid}.{k}: พบ{what} — ตรวจก่อนเผยแพร่")
        n_actions = len(split_lines(merged["actions_today"]))
        if "actions_today" in changed and not 2 <= n_actions <= 3:
            warnings.append(f"{cid}: สิ่งที่ลูกเรือต้องทำวันนี้มี {n_actions} ข้อ (ควร 2–3)")

        status = STATUS_OF[decision]
        if row["status"] == "published" and status != "published":
            warnings.append(f"{cid}: เคยเผยแพร่แล้ว จะถูกถอดออกจากแอป ({decision})")
        if status != row["status"]:
            changed["status"] = status
        if status == "published" and changed:
            changed["reviewed"] = today
        if changed:
            updates[cid] = changed
        if note:
            notes.append((cid, decision, row["title"], note))

    print(f"ไฟล์: {path}")
    print("การตัดสิน: " + " · ".join(f"{k} {v}" for k, v in count.items()))
    for cid, changed in updates.items():
        print(f"\n{cid}")
        for k, v in changed.items():
            print("  {}: {}  →  {}".format(k, *diff_pair(by_id[cid].get(k, ""), v)))
    for label, items in [("ชนกัน", conflicts), ("ผิดพลาด", errors), ("เตือน", warnings)]:
        for item in items:
            print(f"{label}: {item}")

    if not write:
        print(f"\n(ดูอย่างเดียว) จะแก้ {len(updates)} เคส — รันซ้ำโดยเติม --write เพื่อเขียนจริง")
        return
    if not updates and not notes:
        print("\nไม่มีอะไรต้องเขียน")
        return

    if updates:
        backup = REVIEW_DIR / "backup" / f"intl_seed_{datetime.now():%Y%m%d_%H%M%S}.csv"
        backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(SEED, backup)
        for cid, changed in updates.items():
            by_id[cid].update(changed)
        write_seed(fieldnames, seed)
        print(f"\nเขียน {len(updates)} เคสลง {SEED.relative_to(ROOT)} (สำรองเดิมไว้ที่ {backup.relative_to(ROOT)})")
    if notes:
        todo = REVIEW_DIR / f"todo_{today}.md"
        lines = [f"# หมายเหตุจากการตรวจ ({path.name}, {today})", ""]
        for cid, decision, title, note in sorted(notes, key=lambda n: (n[1] != "รอแก้", n[0])):
            lines.append(f"- [ ] **{cid}** ({decision}) {title}")
            lines.extend(f"  {line}" for line in note.split("\n"))
        todo.write_text("\n".join(lines) + "\n", encoding="utf-8")
        print(f"หมายเหตุ {len(notes)} รายการ → {todo.relative_to(ROOT)}")
    print("ต่อไป: python tools/build_data.py --include-draft (หรือไม่ใส่ เพื่อดูเฉพาะ published)")


if __name__ == "__main__":
    main()
