"""ทำสารบัญไฟล์ + ดึงข้อความจากเอกสารนิรภัยเก่า (อ่านอย่างเดียว ไม่แก้ไฟล์ต้นทาง)

ใช้ (Git Bash):
    python -P tools/extract_text.py <out_dir นอก repo> <src_dir> [src_dir ...]

ผลลัพธ์ใน out_dir:
    index.csv        ทุกไฟล์: id, path, ext, size, mtime, status, chars
    text/<id>.txt    ข้อความของ pdf/doc/docx/xls/xlsx/pptx/rtf

รันซ้ำได้ ไฟล์ที่ดึงแล้วจะข้าม ไฟล์ .doc ใช้ antiword (มากับ Git for Windows)
ห้ามชี้ out_dir เข้า repo หรือ OneDrive เพราะข้อความที่ได้คือข้อมูลดิบของหน่วย
"""
import csv
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TEXT_EXT = {'.pdf', '.doc', '.docx', '.xls', '.xlsx', '.pptx', '.rtf'}
SKIP_NAME = re.compile(r'^~\$|^thumbs\.db$', re.I)
PDF_MAX_PAGES = 300
# PDF ภาษาไทยหลายไฟล์เก็บสระบน/วรรณยุกต์เป็น glyph ตำแหน่งพิเศษใน Private Use Area
# (U+F700–F71A) แปลงกลับเป็นอักขระไทยปกติ ไม่งั้นค้นคำว่า "ข้อ" "ได้" ไม่เจอ
THAI_PUA = str.maketrans({
    0xF700: 0x0E10, 0xF701: 0x0E34, 0xF702: 0x0E35, 0xF703: 0x0E36, 0xF704: 0x0E37,
    0xF705: 0x0E48, 0xF706: 0x0E49, 0xF707: 0x0E4A, 0xF708: 0x0E4B, 0xF709: 0x0E4C,
    0xF70A: 0x0E48, 0xF70B: 0x0E49, 0xF70C: 0x0E4A, 0xF70D: 0x0E4B, 0xF70E: 0x0E4C,
    0xF70F: 0x0E0D, 0xF710: 0x0E31, 0xF711: 0x0E4D, 0xF712: 0x0E47, 0xF713: 0x0E48,
    0xF714: 0x0E49, 0xF715: 0x0E4A, 0xF716: 0x0E4B, 0xF717: 0x0E4C, 0xF718: 0x0E38,
    0xF719: 0x0E39, 0xF71A: 0x0E3A,
})


def long(p: str) -> str:
    # path ภาษาไทยลึกๆ ใน E: ยาวเกิน 260 ตัวอักษร ต้องใช้ prefix นี้
    p = os.path.abspath(p)
    return p if p.startswith('\\\\?\\') else '\\\\?\\' + p


def fid(rel: str) -> str:
    return hashlib.sha1(rel.encode('utf-8')).hexdigest()[:12]


def sniff(path: str) -> str:
    with open(path, 'rb') as f:
        head = f.read(8)
    if head.startswith(b'\xd0\xcf\x11\xe0'):
        return 'ole'
    if head.startswith(b'PK'):
        return 'zip'
    if head.startswith(b'{\\rtf'):
        return 'rtf'
    if head.startswith(b'%PDF'):
        return 'pdf'
    return 'other'


def rtf_text(path: str) -> str:
    # แบบหยาบ: ถอด \uNNNN และ \'xx (cp874) แล้วตัดคำสั่ง RTF ทิ้ง
    raw = open(path, 'rb').read().decode('latin-1')
    raw = re.sub(r'\\u(-?\d+)\??', lambda m: chr(int(m.group(1)) % 65536), raw)
    raw = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes([int(m.group(1), 16)]).decode('cp874', 'ignore'), raw)
    raw = re.sub(r'\\[a-z]+-?\d* ?|[{}]', '', raw)
    return raw


def pdf_text(path: str) -> str:
    import fitz
    with fitz.open(path) as d:
        return '\n'.join(d[i].get_text() for i in range(min(len(d), PDF_MAX_PAGES)))


def docx_text(path: str) -> str:
    import docx
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out.append(' | '.join(c.text for c in r.cells))
    return '\n'.join(out)


def doc_text(path: str) -> str:
    # antiword อ่านชื่อไฟล์ภาษาไทยไม่ได้ จึงคัดลอกไปชื่อชั่วคราวก่อน
    with tempfile.TemporaryDirectory() as td:
        tmp = os.path.join(td, 'a.doc')
        shutil.copyfile(path, tmp)
        r = subprocess.run(['antiword', '-m', 'UTF-8.txt', tmp], capture_output=True, timeout=120)
        return r.stdout.decode('utf-8', 'ignore')


def xlsx_text(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f'## {ws.title}')
        for row in ws.iter_rows(values_only=True):
            cells = [str(v) for v in row if v not in (None, '')]
            if cells:
                out.append(' | '.join(cells))
    wb.close()
    return '\n'.join(out)


def xls_text(path: str) -> str:
    import xlrd
    wb = xlrd.open_workbook(path, on_demand=True)
    out = []
    for ws in wb.sheets():
        out.append(f'## {ws.name}')
        for i in range(ws.nrows):
            cells = [str(v) for v in ws.row_values(i) if v not in (None, '')]
            if cells:
                out.append(' | '.join(cells))
    return '\n'.join(out)


def pptx_text(path: str) -> str:
    # อ่านเฉพาะ XML ของสไลด์ ไม่แตกรูป/วิดีโอที่ฝังไว้ (pptx รวมกัน 100+ GB)
    out = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)]
        for n in sorted(names, key=lambda s: int(re.search(r'(\d+)', s.rsplit('/', 1)[1]).group(1))):
            xml = z.read(n).decode('utf-8', 'ignore')
            out.append('\n'.join(re.findall(r'<a:t>([^<]*)</a:t>', xml)))
    return '\n\n'.join(out)


def extract(path: str, ext: str) -> str:
    kind = sniff(path)
    if kind == 'pdf':
        return pdf_text(path)
    if kind == 'rtf':
        return rtf_text(path)
    if ext in ('.doc', '.docx'):
        return doc_text(path) if kind == 'ole' else docx_text(path)
    if ext in ('.xls', '.xlsx'):
        return xls_text(path) if kind == 'ole' else xlsx_text(path)
    if ext == '.pptx':
        return pptx_text(path)
    raise ValueError(f'unknown format ({kind})')


def work(job):
    path, out = job
    try:
        text = extract(path, Path(path).suffix.lower())
        text = re.sub(r'[ \t]+\n', '\n', text.translate(THAI_PUA)).strip()
        with open(out, 'w', encoding='utf-8') as f:
            f.write(text)
        return 'ok' if text else 'empty', len(text)
    except Exception as e:
        return f'error: {type(e).__name__}: {str(e)[:120]}', 0


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    out_dir = Path(sys.argv[1]).resolve()
    for bad in (ROOT, Path.home() / 'OneDrive'):
        if out_dir == bad or bad in out_dir.parents:
            sys.exit(f'out_dir อยู่ใน {bad} — ต้องอยู่นอก repo และนอก OneDrive')
    (out_dir / 'text').mkdir(parents=True, exist_ok=True)

    rows, jobs = [], []
    for src in sys.argv[2:]:
        src_l = long(src)
        for dp, _, fs in os.walk(src_l):
            for name in fs:
                full = os.path.join(dp, name)
                rel = os.path.relpath(full, long(os.path.dirname(os.path.abspath(src)))).replace('\\', '/')
                ext = Path(name).suffix.lower()
                try:
                    st = os.stat(full)
                except OSError:
                    continue
                row = {'id': fid(rel), 'path': rel, 'ext': ext, 'size': st.st_size,
                       'mtime': datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d'),
                       'status': '', 'chars': ''}
                rows.append(row)
                if ext in TEXT_EXT and not SKIP_NAME.search(name) and st.st_size > 0:
                    out = out_dir / 'text' / f"{row['id']}.txt"
                    if out.exists():
                        row['status'], row['chars'] = 'ok', out.stat().st_size
                    else:
                        jobs.append((row, full, str(out)))
    print(f'{len(rows)} ไฟล์, ต้องดึงข้อความ {len(jobs)} ไฟล์', flush=True)
    if not rows:
        sys.exit('ไม่พบไฟล์ใน src')

    done = 0
    with ProcessPoolExecutor(max_workers=max(2, (os.cpu_count() or 4) - 2)) as ex:
        futs = {ex.submit(work, (full, out)): row for row, full, out in jobs}
        for fut in as_completed(futs):
            row = futs[fut]
            try:
                row['status'], row['chars'] = fut.result()
            except Exception as e:
                row['status'] = f'error: {type(e).__name__}'
            done += 1
            if done % 500 == 0:
                print(f'  {done}/{len(jobs)}', flush=True)

    with open(out_dir / 'index.csv', 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    st = {}
    for r in rows:
        if r['status']:
            k = r['status'].split(':')[0]
            st[k] = st.get(k, 0) + 1
    print('เสร็จ:', st)


if __name__ == '__main__':
    main()
